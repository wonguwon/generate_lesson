# 출력/9월_일보의 .xls 5개를 .xlsx로 바꾸고, 원본/영수증 사진 6장을 PDF 하나로 묶는 스크립트
# (3차시 "Codex의 코딩 도구 활용" 결과로 '출력/스크립트'에 남는 스크립트의 예시. 4·5차시 시작 자료에 들어감)
# 필요: Python 3.10 이상, openpyxl, xlrd, Pillow
# 사용: python 형식변환.py <다온유통_영업지원팀 폴더 경로>
import os
import sys

import xlrd
from openpyxl import Workbook
from PIL import Image

BASE = sys.argv[1] if len(sys.argv) > 1 else "."
ILBO = os.path.join(BASE, "출력", "9월_일보")
RECEIPT = os.path.join(BASE, "원본", "영수증")
PDF = os.path.join(BASE, "출력", "영수증_9월.pdf")


def xls_to_xlsx(src, dst):
    book = xlrd.open_workbook(src)
    wb = Workbook()
    wb.remove(wb.active)
    for sh in book.sheets():                      # 시트를 모두 그대로 옮김 (일보, 메모)
        ws = wb.create_sheet(sh.name)
        for r in range(sh.nrows):
            row = []
            for c in range(sh.ncols):
                cell = sh.cell(r, c)
                if cell.ctype == xlrd.XL_CELL_DATE:
                    row.append(xlrd.xldate_as_datetime(cell.value, book.datemode))
                elif cell.ctype == xlrd.XL_CELL_NUMBER and float(cell.value).is_integer():
                    row.append(int(cell.value))
                else:
                    row.append(cell.value)
            ws.append(row)
            if r > 0 and sh.name == "일보":
                ws.cell(row=ws.max_row, column=1).number_format = "yyyy-mm-dd"
    wb.save(dst)


def receipts_to_pdf():
    names = sorted(f for f in os.listdir(RECEIPT) if f.lower().endswith(".jpg"))
    pages = [Image.open(os.path.join(RECEIPT, n)).convert("RGB") for n in names]
    pages[0].save(PDF, "PDF", save_all=True, append_images=pages[1:], resolution=100)
    return len(pages)


if __name__ == "__main__":
    done = 0
    for f in sorted(os.listdir(ILBO)):
        if f.endswith(".xls"):
            src = os.path.join(ILBO, f)
            xls_to_xlsx(src, src + "x")
            os.remove(src)                        # 변환이 끝난 .xls는 지움 ('출력' 쪽만)
            done += 1
    n = receipts_to_pdf()
    print(f"xls → xlsx {done}개, 영수증 PDF {n}쪽")
