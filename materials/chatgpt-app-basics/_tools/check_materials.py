# 만든 실습 자료 파일을 직접 다시 열어 교안 숫자와 대조하는 검산 스크립트
# 사용: python check_materials.py [자료 폴더 경로]   (기본: chatgpt-app-basics)
import datetime as dt
import os
import re
import sys
import zipfile
from collections import defaultdict

import xlrd
from openpyxl import load_workbook

TOOLS = os.path.dirname(os.path.abspath(__file__))
MAT = sys.argv[1] if len(sys.argv) > 1 else os.path.dirname(TOOLS)
P = "다온유통_영업지원팀"
STORES = ["강남점", "송파점", "분당점", "일산점", "인천점", "수원점"]
CANON = {"강남": "강남점", "분당지점": "분당점", "수원 점": "수원점", "송파": "송파점"}
fails = []


def ok(label, got, want):
    mark = "OK " if got == want else "NG "
    if got != want:
        fails.append(label)
    print(f"  {mark}{label}: {got}" + ("" if got == want else f"  (교안: {want})"))


def read(path):
    if path.endswith(".xls"):
        b = xlrd.open_workbook(path); s = b.sheet_by_name("일보")
        rows = []
        for i in range(1, s.nrows):
            v = s.row_values(i); v[0] = xlrd.xldate_as_datetime(v[0], b.datemode).date()
            rows.append([v[0], v[1], v[2], int(v[3]), int(v[4])])
        return b.sheet_names(), s.row_values(0), rows
    wb = load_workbook(path); vals = list(wb["일보"].values)
    rows = [[r[0].date() if isinstance(r[0], dt.datetime) else r[0], *r[1:]] for r in vals[1:]]
    return wb.sheetnames, list(vals[0]), rows


def month_check(d, n_files, raw_rows, raw_qty, rows, qty, amt, per_file=None):
    files = sorted(os.listdir(d))
    ok("파일 수", len(files), n_files)
    allr = []; pf = {}
    for f in files:
        sh, hdr, rs = read(os.path.join(d, f))
        pf[f] = (len(rs), sum(r[3] for r in rs))
        allr += rs
    ok("파일별 데이터 행 합", len(allr), raw_rows)
    ok("파일별 판매수량 합", sum(r[3] for r in allr), raw_qty)
    if per_file:
        bad = [f for f, v in per_file.items() if pf.get(f) != v]
        ok("파일별 행 수·수량 (교안 표)", bad, [])
    seen = set(); clean = []
    for r in allr:
        t = tuple(r)
        if t in seen:
            continue
        seen.add(t)
        if r[3] < 0:
            continue
        clean.append([r[0], CANON.get(r[1], r[1]), r[2], r[3], r[4]])
    ok("정리 후 행 수", len(clean), rows)
    ok("정리 후 판매수량", sum(r[3] for r in clean), qty)
    ok("정리 후 판매금액", sum(r[4] for r in clean), amt)
    return clean


def summarize(clean, weeks):
    ta = sum(r[4] for r in clean)
    st = {s: (sum(r[3] for r in clean if r[1] == s), sum(r[4] for r in clean if r[1] == s)) for s in STORES}
    wk = [(sum(r[3] for r in clean if a <= r[0] <= b), sum(r[4] for r in clean if a <= r[0] <= b)) for a, b in weeks]
    return ta, st, wk


# 9월 교안 표 (지금 이름: 행 수, 판매수량)
PF9 = {"일보 9-1.xls": (36, 747), "일보 9-2.xls": (36, 723), "일보 9-3.xls": (36, 729), "일보 9-4.xls": (36, 740),
       "일보 9-5.xls": (36, 871), "0906_일보.xlsx": (36, 858), "0907_일보.xlsx": (36, 697), "0908_일보.xlsx": (36, 710),
       "0909_일보.xlsx": (37, 749), "0910_일보.xlsx": (36, 768), "0911_일보.xlsx": (36, 731), "12.09 일보.xlsx": (36, 854),
       "9월13일 판매.xlsx": (36, 885), "9월14일 판매.xlsx": (36, 741), "9월15일 판매.xlsx": (36, 737),
       "9월16일 판매.xlsx": (39, 780), "9월17일 판매.xlsx": (36, 730), "9월18일 판매.xlsx": (36, 758),
       "9월19일 판매.xlsx": (36, 879), "일보_20260920.xlsx": (36, 889), "일보_20260921.xlsx": (36, 748),
       "일보_20260922.xlsx": (36, 746), "일보_20260923.xlsx": (36, 714), "일보_20260924.xlsx": (36, 732),
       "일보_20260925.xlsx": (0, 0), "일보_20260926.xlsx": (36, 860), "9월 일보(최종).xlsx": (36, 850),
       "판매일보(9.28).xlsx": (36, 724), "판매일보(9.29).xlsx": (36, 739), "판매일보(9.30).xlsx": (37, 733)}
PF10 = {"일보 10-1.xls": (36, 725), "일보 10-2.xls": (36, 742), "1003_일보.xlsx": (36, 857), "1004_일보.xlsx": (36, 882),
        "1005_일보.xlsx": (36, 728), "1006_일보.xlsx": (36, 733), "1007_일보.xlsx": (36, 720), "1008_일보.xlsx": (36, 750),
        "1009_일보.xlsx": (36, 738), "1010_일보.xlsx": (36, 876), "1011_일보.xlsx": (36, 866),
        "10월12일 판매.xlsx": (36, 724), "10월13일 판매.xlsx": (36, 757), "10월14일 판매.xlsx": (36, 718),
        "10월15일 판매.xlsx": (38, 770), "10월16일 판매.xlsx": (36, 740), "10월17일 판매.xlsx": (36, 863),
        "10월18일 판매.xlsx": (36, 863), "일보_20261019.xlsx": (36, 722), "일보_20261020.xlsx": (37, 701),
        "일보_20261021.xlsx": (36, 742), "일보_20261022.xlsx": (36, 759), "일보_20261023.xlsx": (36, 728),
        "일보_20261024.xlsx": (36, 856), "일보_20261025.xlsx": (36, 878), "판매일보(10.26).xlsx": (36, 751),
        "판매일보(10.27).xlsx": (36, 728), "판매일보(10.28).xlsx": (36, 739), "판매일보(10.29).xlsx": (36, 744),
        "판매일보(10.30).xlsx": (36, 731), "10월 일보(최종).xlsx": (36, 887)}
D = dt.date
W9 = [(D(2026, 9, 1), D(2026, 9, 6)), (D(2026, 9, 7), D(2026, 9, 13)), (D(2026, 9, 14), D(2026, 9, 20)),
      (D(2026, 9, 21), D(2026, 9, 27)), (D(2026, 9, 28), D(2026, 9, 30))]
W10 = [(D(2026, 10, 1), D(2026, 10, 4)), (D(2026, 10, 5), D(2026, 10, 11)), (D(2026, 10, 12), D(2026, 10, 18)),
       (D(2026, 10, 19), D(2026, 10, 25)), (D(2026, 10, 26), D(2026, 10, 31))]

print("[1차시] 처음 파일 8개")
base = os.path.join(MAT, "01_codex-agent", P)
ok("파일 수", len(os.listdir(base)), 8)
for f, n in [("거래처목록.xlsx", 18), ("8월_지점별_매출.xlsx", 6), ("9월_지점별_매출.xlsx", 6), ("반품접수_9월.xlsx", 23)]:
    ws = load_workbook(os.path.join(base, f)).active
    ok(f"{f} 데이터 행", ws.max_row - 1, n)
ws = load_workbook(os.path.join(base, "반품접수_9월.xlsx")).active
agg = defaultdict(lambda: [0, 0])
for r in list(ws.values)[1:]:
    agg[r[4]][0] += 1; agg[r[4]][1] += r[3]
ok("반품 사유별 (건수, 수량)", dict(agg), {"포장 파손": [12, 72], "수량 착오": [5, 38], "유통기한 임박": [3, 25], "오배송": [3, 17]})
s9 = {r[0]: r[1] for r in list(load_workbook(os.path.join(base, "9월_지점별_매출.xlsx")).active.values)[1:]}

print("[3차시] 9월 일보 30개 (03_batch-files/9월_일보)")
c9 = month_check(os.path.join(MAT, "03_batch-files", "9월_일보"), 30, 1049, 22422, 1044, 22357, 202883040, PF9)
ok("영수증 사진 수", len(os.listdir(os.path.join(MAT, "03_batch-files", "영수증"))), 6)
ta, st, wk = summarize(c9, W9)
ok("지점별 (수량, 금액)", st, {"강남점": (4297, 39353380), "송파점": (3741, 33837000), "분당점": (4000, 36455060),
                          "일산점": (3200, 28913100), "인천점": (3457, 31264120), "수원점": (3662, 33060380)})
ok("주차별 (수량, 금액)", wk, [(4668, 42364900), (5398, 48978400), (5443, 49324200), (4650, 42233500), (2198, 19982040)])
ok("매출비중 (%)", [round(st[s][1] / ta * 100, 1) for s in STORES], [19.4, 16.7, 18.0, 14.3, 15.4, 16.3])
ok("일보 금액 < 전 품목 매출(만원) - 모든 지점", all(st[s][1] < s9[s] * 10000 for s in STORES), True)
dq = defaultdict(int); da = defaultdict(int)
for r in c9:
    dq[r[0]] += r[3]; da[r[0]] += r[4]
ok("판매수량 최고일", (max(dq, key=dq.get), max(dq.values())), (D(2026, 9, 20), 889))
ok("판매수량 최저일", (min(dq, key=dq.get), min(dq.values())), (D(2026, 9, 7), 697))
ok("판매금액 최고일", (max(da, key=da.get), max(da.values())), (D(2026, 9, 13), 8096400))
ok("판매 기록이 있는 날", len(dq), 29)
first = sorted(c9, key=lambda r: (r[0], STORES.index(r[1])))[0]
ok("첫 데이터 행 (9/1 강남점 생수 수량)", (first[1], first[2], first[3]), ("강남점", "생수 2L(6입)", 29))

print("[4차시 시작] 출력/9월_일보 · 9월_일보통합.xlsx")
p4 = os.path.join(MAT, "04_outputs-edit", P)
od = os.path.join(p4, "출력", "9월_일보")
names = sorted(os.listdir(od))
ok("출력/9월_일보 파일 수와 이름", names, [f"일보_2026-09-{d:02d}.xlsx" for d in range(1, 31)])
ok("일보_2026-09-03.xlsx 시트", load_workbook(os.path.join(od, "일보_2026-09-03.xlsx")).sheetnames, ["일보", "메모"])
ok("일보_2026-09-01.xlsx 머리글", list(next(load_workbook(os.path.join(od, "일보_2026-09-01.xlsx"))["일보"].values)),
   ["날짜", "매장", "품목", "수량", "판매금액"])
ws = load_workbook(os.path.join(p4, "출력", "9월_일보통합.xlsx"))["통합"]
vals = list(ws.values)
ok("통합 머리글", list(vals[0]), ["날짜", "지점", "품목", "판매수량", "판매금액", "출처파일"])
ok("통합 행 수", len(vals) - 1, 1044)
ok("통합 판매수량", sum(r[3] for r in vals[1:]), 22357)
ok("통합 판매금액", sum(r[4] for r in vals[1:]), 202883040)
ok("통합 지점명", sorted({r[1] for r in vals[1:]}), sorted(STORES))
ok("이름변경목록 행 수", sum(1 for l in open(os.path.join(p4, "출력", "이름변경목록.md"), encoding="utf-8") if l.startswith("| 9/")), 30)
pdf = open(os.path.join(p4, "출력", "영수증_9월.pdf"), "rb").read()
ok("영수증 PDF 쪽수", len(re.findall(rb"/Type\s*/Page[^s]", pdf)), 6)
rec = open(os.path.join(p4, "출력", "전처리기록.md"), encoding="utf-8").read()
for k in ["1,044행", "22,357", "202,883,040원", "**18**", "9/3 강남점 냉장 진열대 점검으로 오후 2시간 판매 중단", "일보_2026-09-25.xlsx"]:
    ok(f"전처리기록에 '{k}'", k in rec, True)

print("[5차시 시작] 4차시 결과물")
p5 = os.path.join(MAT, "05_agents-skills", P)
wb = load_workbook(os.path.join(p5, "출력", "9월_판매요약.xlsx"), data_only=True)
ok("요약 시트", wb.sheetnames, ["지점별", "주차별", "데이터"])
ok("지점별 강남점 (값)", [wb["지점별"]["B2"].value, wb["지점별"]["C2"].value], [4297, 39353380])
ok("지점별 합계 (값)", [wb["지점별"]["B8"].value, wb["지점별"]["C8"].value, round(wb["지점별"]["D8"].value, 3)], [22357, 202883040, 1.0])
ok("주차별 판매금액 (값)", [wb["주차별"].cell(row=r, column=5).value for r in range(2, 8)],
   [42364900, 48978400, 49324200, 42233500, 19982040, 202883040])
wf = load_workbook(os.path.join(p5, "출력", "9월_판매요약.xlsx"))
ok("강남점 판매금액 칸이 수식", str(wf["지점별"]["C2"].value).startswith("=SUMIFS("), True)
ok("데이터 시트 첫 행 판매수량", wf["데이터"]["D2"].value, 29)


def zip_text(path, pattern):
    z = zipfile.ZipFile(path)
    return "".join(re.sub(r"<[^>]+>", "", z.read(n).decode("utf-8")) for n in z.namelist() if re.match(pattern, n))


docx = zip_text(os.path.join(p5, "출력", "주간보고_9월5주차.docx"), r"word/document.xml")
for k in ["영업지원팀 주간보고 (9월 5주차)", "22,357", "202,883,040", "29일", "39,353,380", "49,324,200",
          "9월 반품 23건(152개) 중 포장 파손이 12건(72개)으로 가장 많음", "10월 일보 정리 규칙을 문서로 남기기"]:
    ok(f"주간보고에 '{k}'", k in docx, True)
z = zipfile.ZipFile(os.path.join(p5, "출력", "9월_판매실적.pptx"))
ok("발표자료 슬라이드 수", len([n for n in z.namelist() if re.match(r"ppt/slides/slide\d+\.xml$", n)]), 5)
charts = [n for n in z.namelist() if re.match(r"ppt/charts/chart\d+\.xml$", n)]
ok("편집 가능한 차트 수", len(charts), 2)
ok("차트 원본 데이터(엑셀) 포함", len([n for n in z.namelist() if n.startswith("ppt/embeddings/")]), 2)
ctext = "".join(z.read(n).decode("utf-8") for n in charts)
for k in ["39353380", "28913100", "19982040", "1주차 9/1~9/6", "5주차 9/28~9/30"]:
    ok(f"차트 데이터에 '{k}'", k in ctext, True)
html = open(os.path.join(p5, "출력", "9월_대시보드.html"), encoding="utf-8").read()
urls = [u for u in re.findall(r"https?://[^\s\"'<>)]+", html) if u != "http://www.w3.org/2000/svg"]
ok("대시보드 외부 주소 (svg xmlns 제외)", urls, [])
ok("대시보드 제목", re.search(r"<h1>(.*?)</h1>", html).group(1), "다온유통 9월 판매 대시보드")
ok("대시보드 점 수", html.count("<circle"), 29)
ok("대시보드 '추석 휴무'", "추석 휴무" in html, True)
ok("출력 폴더 목록 (5차시 시작)", sorted(os.listdir(os.path.join(p5, "출력"))),
   sorted(["폴더요약.md", "반품요약.md", "9월_일보", "이름변경목록.md", "스크립트", "영수증_9월.pdf", "9월_일보통합.xlsx",
           "전처리기록.md", "기대값.md", "9월_판매요약.xlsx", "주간보고_9월5주차.docx", "9월_판매실적.pptx",
           "9월_대시보드.html", "대조결과.md"]))

print("[5차시] 10월 일보 31개 (05_agents-skills/10월_일보)")
c10 = month_check(os.path.join(MAT, "05_agents-skills", "10월_일보"), 31, 1119, 24018, 1116, 23966, 217265920, PF10)
ta, st, wk = summarize(c10, W10)
ok("10월 지점별 (수량, 금액)", st, {"강남점": (4572, 41977640), "송파점": (4083, 36933800), "분당점": (4283, 38857940),
                              "일산점": (3500, 31514440), "인천점": (3688, 33138380), "수원점": (3840, 34843720)})
ok("10월 주차별 (수량, 금액)", wk, [(3206, 28904260), (5411, 48973980), (5380, 48704220), (5389, 49002940), (4580, 41680520)])
ok("10월 천 원 지점별 합", sum(round(st[s][1] / 1000) for s in STORES), round(ta / 1000))
ok("10월 매출비중 (%)", [round(st[s][1] / ta * 100, 1) for s in STORES], [19.3, 17.0, 17.9, 14.5, 15.3, 16.0])

print()
print("결과:", "모두 일치" if not fails else f"불일치 {len(fails)}건 - {fails}")
