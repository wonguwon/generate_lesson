# 다온유통 9월 일보 30개 + 영수증 사진 6장 만들기 (자료 제작용)
# 필요: Python 3.10 이상, openpyxl, xlwt, Pillow  (pip install openpyxl xlwt pillow)
# 사용: python make_ilbo.py <다온유통_영업지원팀 폴더 경로>
# 결과: '원본/9월_일보'(30개)와 '원본/영수증'(6장)이 생깁니다.
# 영수증 글꼴은 Windows의 맑은 고딕을 씁니다. 다른 OS에서는 한글이 깨질 수 있습니다.
import datetime as dt
import os
import random
import sys
from openpyxl import Workbook
import xlwt
from PIL import Image, ImageDraw, ImageFont

BASE = sys.argv[1] if len(sys.argv) > 1 else "."
ILBO = os.path.join(BASE, "원본", "9월_일보")
RECEIPT = os.path.join(BASE, "원본", "영수증")

STORES = [("강남점", 12), ("송파점", 9), ("분당점", 10),
          ("일산점", 6), ("인천점", 7), ("수원점", 8)]
ITEMS = [("생수 2L(6입)", 14, 5900), ("라면 5입", 10, 4500),
         ("즉석밥 12입", 6, 13800), ("참치캔 3입", 8, 7980),
         ("식용유 1.8L", 3, 8900), ("두루마리 휴지 30롤", 5, 15900)]
PRICE = {name: price for name, _, price in ITEMS}

rng = random.Random(2126)            # 같은 숫자가 나오도록 씨앗값 고정
QTY = {}
for day in range(1, 31):
    weekend = 4 if dt.date(2026, 9, day).weekday() >= 5 else 0   # 토·일 +4
    for si in range(6):
        for pi in range(6):
            QTY[day, si, pi] = STORES[si][1] + ITEMS[pi][1] + weekend + rng.randint(0, 8)


def file_name(day):
    if day <= 5:
        return f"일보 9-{day}.xls"
    if day <= 11:
        return f"09{day:02d}_일보.xlsx"
    if day == 12:
        return "12.09 일보.xlsx"
    if day <= 19:
        return f"9월{day}일 판매.xlsx"
    if day <= 26:
        return f"일보_202609{day}.xlsx"
    if day == 27:
        return "9월 일보(최종).xlsx"
    return f"판매일보(9.{day}).xlsx"


RENAME_STORE = {14: ("강남점", "강남"), 22: ("분당점", "분당지점"), 29: ("수원점", "수원 점")}


def rows_for(day):
    if day == 25:                      # 추석 당일: 머리글만 있는 빈 파일
        return []
    date = dt.date(2026, 9, day)
    rows = []
    for si, (store, _) in enumerate(STORES):
        name = store
        if day in RENAME_STORE and RENAME_STORE[day][0] == store:
            name = RENAME_STORE[day][1]    # 지점명 표기 흔들림
        for pi, (item, _, price) in enumerate(ITEMS):
            q = QTY[day, si, pi]
            rows.append([date, name, item, q, q * price])
    if day == 16:                      # 중복 행 3개 (송파점 앞 3품목을 한 번 더)
        rows += [list(r) for r in rows[6:9]]
    if day == 9:                       # 음수 수량 1행
        rows.append([date, "강남점", "생수 2L(6입)", -4, -4 * PRICE["생수 2L(6입)"]])
    if day == 30:                      # 음수 수량 1행
        rows.append([date, "인천점", "식용유 1.8L", -2, -2 * PRICE["식용유 1.8L"]])
    return rows


def header_for(day):
    if day <= 5:                       # 예전 양식: 열 이름이 다름
        return ["날짜", "매장", "품목", "수량", "판매금액"]
    return ["날짜", "지점", "품목", "판매수량", "판매금액"]


def write_xlsx(path, header, rows):
    wb = Workbook()
    ws = wb.active
    ws.title = "일보"
    ws.append(header)
    for r in rows:
        ws.append(r)
        ws.cell(row=ws.max_row, column=1).number_format = "yyyy-mm-dd"
    wb.save(path)


def write_xls(path, header, rows, memo=None):
    wb = xlwt.Workbook()
    ws = wb.add_sheet("일보")
    date_style = xlwt.easyxf(num_format_str="yyyy-mm-dd")
    for c, h in enumerate(header):
        ws.write(0, c, h)
    for r, row in enumerate(rows, start=1):
        ws.write(r, 0, row[0], date_style)
        for c in range(1, 5):
            ws.write(r, c, row[c])
    if memo:                           # 시트가 2개인 파일
        ms = wb.add_sheet("메모")
        ms.write(0, 0, memo)
    wb.save(path)


RECEIPTS = [  # (파일 이름, 날짜, 내용, 금액)
    ("IMG_20260903_091512.jpg", "2026-09-03", "택시 - 강남점 방문", 12400),
    ("IMG_20260908_123040.jpg", "2026-09-08", "식대 - 송파점 점검", 18000),
    ("IMG_20260911_180205.jpg", "2026-09-11", "주차 - 분당점 방문", 6000),
    ("IMG_20260917_132210.jpg", "2026-09-17", "택시 - 일산점 방문", 23700),
    ("IMG_20260922_094530.jpg", "2026-09-22", "식대 - 인천점 점검", 15500),
    ("IMG_20260929_171801.jpg", "2026-09-29", "주차 - 수원점 방문", 4000),
]


def write_receipt(path, n, date, what, amount):
    img = Image.new("RGB", (600, 900), "white")
    d = ImageDraw.Draw(img)
    try:
        font = ImageFont.truetype("malgun.ttf", 34)   # Windows 맑은 고딕
    except OSError:
        font = ImageFont.load_default()
    lines = ["영 수 증 (연습용)", f"번호: {n}", f"날짜: {date}", what, f"금액: {amount:,}원"]
    for i, t in enumerate(lines):
        d.text((50, 80 + i * 110), t, fill="black", font=font)
    img.save(path, "JPEG")


if __name__ == "__main__":
    os.makedirs(ILBO, exist_ok=True)
    os.makedirs(RECEIPT, exist_ok=True)
    for day in range(1, 31):
        path = os.path.join(ILBO, file_name(day))
        if day <= 5:
            memo = "9/3 강남점 냉장 진열대 점검으로 오후 2시간 판매 중단" if day == 3 else None
            write_xls(path, header_for(day), rows_for(day), memo)
        else:
            write_xlsx(path, header_for(day), rows_for(day))
    for n, (name, date, what, amount) in enumerate(RECEIPTS, start=1):
        write_receipt(os.path.join(RECEIPT, name), n, date, what, amount)
    print("완료: 일보 30개, 영수증 6장")
