# 다온유통 10월 일보 31개 만들기 (자료 제작용)
# 필요: Python 3.10 이상, openpyxl, xlwt  (pip install openpyxl xlwt)
# 사용: python make_ilbo_10.py <다온유통_영업지원팀 폴더 경로>
# 결과: '원본/10월_일보'(31개)가 생깁니다.
import datetime as dt
import os
import random
import sys
from openpyxl import Workbook
import xlwt

BASE = sys.argv[1] if len(sys.argv) > 1 else "."
ILBO = os.path.join(BASE, "원본", "10월_일보")

STORES = [("강남점", 12), ("송파점", 9), ("분당점", 10),
          ("일산점", 6), ("인천점", 7), ("수원점", 8)]
ITEMS = [("생수 2L(6입)", 14, 5900), ("라면 5입", 10, 4500),
         ("즉석밥 12입", 6, 13800), ("참치캔 3입", 8, 7980),
         ("식용유 1.8L", 3, 8900), ("두루마리 휴지 30롤", 5, 15900)]
PRICE = {name: price for name, _, price in ITEMS}

rng = random.Random(1031)            # 같은 숫자가 나오도록 씨앗값 고정
QTY = {}
for day in range(1, 32):
    weekend = 4 if dt.date(2026, 10, day).weekday() >= 5 else 0   # 토·일 +4
    for si in range(6):
        for pi in range(6):
            QTY[day, si, pi] = STORES[si][1] + ITEMS[pi][1] + weekend + rng.randint(0, 8)


def file_name(day):
    if day <= 2:
        return f"일보 10-{day}.xls"
    if day <= 11:
        return f"10{day:02d}_일보.xlsx"
    if day <= 18:
        return f"10월{day}일 판매.xlsx"
    if day <= 25:
        return f"일보_202610{day}.xlsx"
    if day == 31:
        return "10월 일보(최종).xlsx"
    return f"판매일보(10.{day}).xlsx"


def rows_for(day):
    date = dt.date(2026, 10, day)
    rows = []
    for si, (store, _) in enumerate(STORES):
        name = "송파" if (day == 12 and store == "송파점") else store   # 지점명 표기 흔들림
        for pi, (item, _, price) in enumerate(ITEMS):
            q = QTY[day, si, pi]
            rows.append([date, name, item, q, q * price])
    if day == 15:                      # 중복 행 2개 (분당점 앞 2품목을 한 번 더)
        rows += [list(r) for r in rows[12:14]]
    if day == 20:                      # 음수 수량 1행
        rows.append([date, "수원점", "라면 5입", -3, -3 * PRICE["라면 5입"]])
    return rows


def header_for(day):
    if day <= 2:                       # 예전 양식: 열 이름이 다름
        return ["날짜", "매장", "품목", "수량", "판매금액"]
    return ["날짜", "지점", "품목", "판매수량", "판매금액"]


def write_xlsx(path, header, rows):
    wb = Workbook()
    ws = wb.active
    ws.title = "일보"
    ws.append(header)
    for r in rows:
        ws.append(r)
        ws.cell(row=ws.max_row, column=1).number_format = "yyyy-mm-dd"
    wb.save(path)


def write_xls(path, header, rows):
    wb = xlwt.Workbook()
    ws = wb.add_sheet("일보")
    date_style = xlwt.easyxf(num_format_str="yyyy-mm-dd")
    for c, h in enumerate(header):
        ws.write(0, c, h)
    for r, row in enumerate(rows, start=1):
        ws.write(r, 0, row[0], date_style)
        for c in range(1, 5):
            ws.write(r, c, row[c])
    wb.save(path)


if __name__ == "__main__":
    os.makedirs(ILBO, exist_ok=True)
    for day in range(1, 32):
        path = os.path.join(ILBO, file_name(day))
        if day <= 2:
            write_xls(path, header_for(day), rows_for(day))
        else:
            write_xlsx(path, header_for(day), rows_for(day))
    print("완료: 10월 일보 31개")
