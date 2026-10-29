"""하루담리빙 주간 실적 보고서(.docx)와 양식(.docx) 생성.

보고서 요청의 규칙을 따른다: 목차 5개(제목 그대로), 비중은 소수 첫째 자리,
금액은 원 단위에 천 단위 쉼표, 파일로 확인되는 사실만 쓰고 추측·전망은 쓰지 않는다.
"""
from docx.shared import Pt

import haru_data as H
import writers as W

SECTIONS = ["1. 요약", "2. 지점별 실적", "3. 상품별 실적", "4. 일자별 실적", "5. 데이터 기준"]


def won(n):
    return f"{n:,}"


def report_lines(rows):
    """요약 3줄에 들어갈 사실 (매출액 1위 지점, 상품 1위(매출액·수량), 매출 최고일)."""
    S = H.summarize(rows)
    tot = S["sales"]
    b1, bv = S["branch"][0]
    p_sales, pv = S["product"][0]
    p_qty, qv = max(S["product"], key=lambda kv: kv[1][1])
    d1, dv = max(S["date"], key=lambda kv: kv[1][2])
    lines = [f"매출액 1위 지점은 {b1}({won(bv[2])}원, 비중 {H.pct(bv[2], tot)}%)입니다."]
    if p_sales == p_qty:
        lines.append(f"상품 1위는 매출액({won(pv[2])}원)과 수량({qv[1]}개) 모두 {p_sales}입니다.")
    else:
        lines.append(f"상품 1위는 매출액 기준 {p_sales}({won(pv[2])}원), 수량 기준 {p_qty}({qv[1]}개)입니다.")
    lines.append(f"매출액이 가장 많은 날은 {H.day_label(d1)} {won(dv[2])}원입니다.")
    return S, lines


def build_report(rows, week, period, source_name, path):
    S, lines = report_lines(rows)
    tot = S["sales"]
    doc = W.new_doc()
    doc.add_heading(f"하루담리빙 {week}주차 실적 보고", level=0)
    doc.add_paragraph(f"기간: {period}\n읽는 사람: 영업팀장")

    doc.add_heading(SECTIONS[0], level=1)
    doc.add_paragraph(f"전체 {S['count']}건, 매출액 합계 {won(tot)}원 (수량 {S['qty']}개)")
    for ln in lines:
        doc.add_paragraph(ln, style="List Bullet")

    doc.add_heading(SECTIONS[1], level=1)
    W.doc_table(doc, ["지점", "건수", "수량", "매출액(원)", "비중(%)"],
                [(b, v[0], v[1], won(v[2]), f"{H.pct(v[2], tot)}%") for b, v in S["branch"]]
                + [("전체", S["count"], S["qty"], won(tot), "100.0%")], right_cols=(1, 2, 3, 4))

    doc.add_heading(SECTIONS[2], level=1)
    W.doc_table(doc, ["상품명", "수량", "매출액(원)", "비중(%)"],
                [(p, v[1], won(v[2]), f"{H.pct(v[2], tot)}%") for p, v in S["product"]]
                + [("전체", S["qty"], won(tot), "100.0%")], right_cols=(1, 2, 3))

    doc.add_heading(SECTIONS[3], level=1)
    W.doc_table(doc, ["판매일", "건수", "매출액(원)"],
                [(H.day_label(d), v[0], won(v[2])) for d, v in S["date"]]
                + [("전체", S["count"], won(tot))], right_cols=(1, 2))

    doc.add_heading(SECTIONS[4], level=1)
    for t in (f"원본 파일: {source_name}", f"기간: {period}", f"건수: {S['count']}건", f"매출액 합계: {won(tot)}원"):
        doc.add_paragraph(t, style="List Bullet")
    W.save_doc(doc, path)
    return S, lines


TEMPLATE_TEXT = """하루담리빙 ○주차 실적 보고
기간: ○○○○-○○-○○(월) ~ ○○○○-○○-○○(금)
읽는 사람: 영업팀장

1. 요약
- 전체 건수와 매출액 합계
- 핵심 내용 3줄

2. 지점별 실적
표: 지점, 건수, 수량, 매출액(원), 비중(%) / 매출액이 큰 순서

3. 상품별 실적
표: 상품명, 수량, 매출액(원), 비중(%) / 매출액이 큰 순서

4. 일자별 실적
표: 판매일, 건수, 매출액(원) / 날짜 순서

5. 데이터 기준
원본 파일 이름, 기간, 건수, 매출액 합계"""


def build_template(path):
    """교안 [자료] 하루담리빙_주간보고_양식.docx 블록 그대로 (숫자 없음)."""
    lines = TEMPLATE_TEXT.split("\n")
    doc = W.new_doc()
    doc.add_heading(lines[0], level=0)
    for ln in lines[1:]:
        if not ln:
            continue
        if ln[:2] in ("1.", "2.", "3.", "4.", "5."):
            doc.add_heading(ln, level=1)
        else:
            doc.add_paragraph(ln)
    W.save_doc(doc, path)
