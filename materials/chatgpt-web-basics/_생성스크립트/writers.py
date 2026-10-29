"""실습 파일 공통 출력 도우미 (xlsx, docx, txt, csv)."""
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Pt
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

HEAD = Font(bold=True, color="FFFFFF")
HEAD_FILL = PatternFill("solid", fgColor="1F4E78")
MARK_FILL = PatternFill("solid", fgColor="F8CBAD")
WRAP = Alignment(wrap_text=True, vertical="top")
KFONT = "맑은 고딕"


def write_text(path, text):
    """Windows 메모장·엑셀에서 한글이 깨지지 않게 UTF-8(BOM), CRLF로 저장."""
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_bytes(text.replace("\r\n", "\n").replace("\n", "\r\n").encode("utf-8-sig"))


def add_sheet(wb, title, header, rows, widths=None, money_cols=(), first=False, notes=()):
    ws = wb.active if first else wb.create_sheet()
    ws.title = title
    ws.append(header)
    for c in ws[1]:
        c.font, c.fill, c.alignment = HEAD, HEAD_FILL, WRAP
    for r in rows:
        ws.append(list(r))
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.alignment = WRAP
            if c.column in money_cols and isinstance(c.value, (int, float)):
                c.number_format = "#,##0"
    widths = widths or [14] * len(header)
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "A2"
    if notes:
        ws.append([])
        for n in notes:
            ws.append([n])
    return ws


def save_data_xlsx(path, header, rows, sheet="데이터", money_cols=(6,)):
    """실습 데이터 파일: 시트 하나, 머리글 1행 + 데이터."""
    wb = Workbook()
    add_sheet(wb, sheet, header, rows, widths=[10, 12, 12, 10, 8, 12], money_cols=money_cols, first=True)
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)


def new_doc():
    doc = Document()
    for name in ("Normal", "Title", "Heading 1", "Heading 2"):
        st = doc.styles[name]
        st.font.name = KFONT
        st.element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), KFONT)
    doc.styles["Normal"].font.size = Pt(10.5)
    return doc


def doc_table(doc, header, rows, right_cols=()):
    t = doc.add_table(rows=1, cols=len(header))
    t.style = "Table Grid"
    for i, h in enumerate(header):
        cell = t.rows[0].cells[i]
        cell.text = h
        cell.paragraphs[0].runs[0].bold = True
    for r in rows:
        cells = t.add_row().cells
        for i, v in enumerate(r):
            cells[i].text = str(v)
            if i in right_cols:
                cells[i].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.RIGHT
    return t


def save_doc(doc, path):
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(path)
