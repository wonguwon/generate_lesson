"""2교시 'Claude 에이전트' 실습 파일 생성. usage: py -3.14 -X utf8 make_02.py <claude-app-basics 폴더>

학습자용/모아상사_실습_2 에 교안 📁 명세의 파일 20개(하위 폴더 없음)를 만든다.
"""
import sys
from collections import Counter
from pathlib import Path

from pptx import Presentation
from pptx.util import Pt as PPt, Cm

import moa_data as M
import writers as W

ROOT = Path(sys.argv[1]) / "02_Claude_에이전트"
L, T = ROOT / "학습자용", ROOT / "강사용"
F = L / "모아상사_실습_2"
EX = T / "예시_결과"

# ---------------- 교안 정답과 대조 ----------------
assert len(M.FILES_2) == 20 and len(set(M.FILES_2)) == 20
ext = Counter(n.rsplit(".", 1)[1] for n in M.FILES_2)
assert [(e, ext[e]) for e in ("docx", "xlsx", "pdf", "jpg", "pptx")] == [
    ("docx", 7), ("xlsx", 5), ("pdf", 4), ("jpg", 3), ("pptx", 1)]
cls = {n: M.classify(n) for n in M.FILES_2}
cc = Counter(cls.values())
assert [(k, cc[k]) for k in ("회의록", "실적", "영수증", "제안서", "기타")] == [
    ("회의록", 6), ("실적", 5), ("영수증", 4), ("제안서", 3), ("기타", 2)]
assert sorted(n for n, c in cls.items() if c == "기타") == ["비상연락망.docx", "사무실_자리배치도.pptx"]
assert sorted(n for n, c in cls.items() if c == "영수증" and n.endswith(".pdf")) == ["영수증_사무용품_2026-10-22.pdf"]
meet_dates = sorted(n[5:15] for n, c in cls.items() if c == "회의록")
assert meet_dates == ["2026-09-21", "2026-09-28", "2026-10-05", "2026-10-12", "2026-10-19", "2026-10-26"]
assert all(M.wd(M.date.fromisoformat(d)) == "월" for d in meet_dates)
assert sorted(meet_dates[:-1]) == sorted(M.PAST_MEETINGS)

totals = {p[0]: M.proposal_total(p) for p in M.PROPOSALS}
assert totals == {"모락오피스": 6660000, "가람테크렌탈": 7776000, "온새미사무기": 5016000}
assert min(M.PROPOSALS, key=lambda p: p[1])[0] == "가람테크렌탈"
assert min(M.PROPOSALS, key=lambda p: M.proposal_total(p))[0] == "온새미사무기"
assert max(M.PROPOSALS, key=lambda p: M.proposal_total(p))[0] == "가람테크렌탈"
by_install = sorted(M.PROPOSALS, key=lambda p: p[3])
assert [(p[0], p[3]) for p in by_install] == [("온새미사무기", 3), ("모락오피스", 7), ("가람테크렌탈", 10)]

# 에이전트와 작업 ④: 안내문 3개 (김지현 팀장은 맡은 할 일 없음)
owners = sorted({n for _, n, _ in M.TODO_1026})
assert owners == ["박민수", "이서연", "최우진"] and "김지현" not in owners
# 실적 파일: 9월 파일은 3교시 명세와 같은 모양·숫자, 8월 실적은 3교시 명세의 8월 실적 열
assert M.month_rows(9) == M.sep_rows() and M.month_header(9) == M.SEP_HEADER
assert [r[2] for r in M.month_rows(8)] == [4180, 3450, 2240]
assert M.month_rows(8)[0][3] == M.MONTH_ACTUAL[7]["영업1팀"]
for k, rows in M.PAST_MEETINGS.items():
    for _, body in rows:
        M.check_weekdays(body, k)
        assert "설치" not in body  # "설치"는 제안서에서만 찾게


# ---------------- 파일 내용 만들기 (쓰기 전에 PDF 글자 추출로 확인) ----------------
def proposal_pdf(p):
    name, monthly, months, days, toner = p
    toner_line = "토너: 포함 (월 임대료에 토너 포함)" if toner == "토너 포함" else "토너: 별도 (월 임대료에 토너 별도)"
    lines = [
        ("복합기 임대 제안서", 20, True), ("", 12, False),
        ("수신: 모아상사 경영지원팀 (실습용 가상 회사)", 11, False),
        (f"제안사: {name} (실습용 가상 거래처)", 11, False), ("", 12, False),
        (f"제안 내용: {M.PROPOSAL_ITEM}", 12, True),
        (f"월 임대료: {monthly:,}원 (부가세 포함)", 12, False),
        (f"계약 기간: {months}개월", 12, False),
        (f"계약 기간 총액: {monthly * months:,}원 ({monthly:,}원 × {months}개월, 부가세 포함)", 12, False),
        (f"설치: 계약 후 {days}영업일", 12, False),
        (toner_line, 12, False), ("", 12, False),
        ("위 조건으로 제안드립니다. 금액은 모두 부가세 포함입니다.", 11, False),
    ]
    data = W.pdf_bytes(f"제안서_{name}", lines, footer="이 문서는 실습용 가상 문서입니다.")
    txt = W.pdf_text(data)
    for must in (name, M.PROPOSAL_ITEM, f"{monthly:,}원", f"{months}개월", f"{monthly * months:,}원",
                 f"계약 후 {days}영업일", toner, "부가세 포함", "설치", "실습용 가상"):
        assert must in txt, f"제안서_{name}.pdf 에서 '{must}'를 찾지 못함:\n{txt}"
    assert txt.count("설치") == 1
    return data


def receipt_pdf(fname):
    store, d, tm, items = M.RECEIPTS[fname]
    total = sum(q * u for _, q, u in items)
    lines = [("영 수 증", 20, True), ("", 10, False), (f"상호: {store}", 11, False),
             (f"일시: {M.label(d)} {tm}", 11, False), ("", 10, False),
             ("품목 / 수량 / 단가 / 금액", 11, True)]
    lines += [(f"{n} / {q}개 / {u:,}원 / {q * u:,}원", 11, False) for n, q, u in items]
    lines += [("", 10, False), (f"합계: {total:,}원 (부가세 포함)", 13, True), ("결제: 법인카드", 11, False)]
    data = W.pdf_bytes(fname[:-4], lines, footer="이 문서는 실습용 가상 영수증입니다.")
    txt = W.pdf_text(data)
    assert f"{total:,}원" in txt and store in txt and "영 수 증" in txt
    return data, total


proposal_data = {p[0]: proposal_pdf(p) for p in M.PROPOSALS}
receipt_data, receipt_total = receipt_pdf("영수증_사무용품_2026-10-22.pdf")
assert receipt_total == 42700


def meeting_text(day, agendas, next_day):
    d = M.date.fromisoformat(day)
    lines = ["[회의록] 모아상사 경영지원팀 주간회의 (실습용 가상 회사)",
             f"일시: {d.year}년 {d.month}월 {d.day}일({M.wd(d)}) 오전 10시",
             "참석: 김지현 팀장, 박민수 대리, 이서연 주임, 최우진 사원", ""]
    for i, (title, body) in enumerate(agendas, start=1):
        lines += [f"{i}. {title}", f"- {body}", ""]
    nd = M.date.fromisoformat(next_day)
    lines.append(f"다음 회의: {nd.year}년 {nd.month}월 {nd.day}일({M.wd(nd)}) 오전 10시")
    return "\n".join(lines) + "\n"


past_texts = {}
days = sorted(M.PAST_MEETINGS) + ["2026-10-26"]
for i, day in enumerate(days[:-1]):
    past_texts[day] = meeting_text(day, M.PAST_MEETINGS[day], days[i + 1])
    M.check_weekdays(past_texts[day], day)

PROMPTS = """Claude 에이전트 - 실습 요청문 모음 (2교시)
교안의 순서대로 적었습니다. 필요한 블록만 복사해 Claude 입력창에 붙여넣습니다.
실습 폴더는 함께 받은 모아상사_실습_2 (바탕화면에 풀기)입니다.

==================================================
[에이전트(Agent)와 작업] ① 작업 맡기기 - 새 대화
괄호 자리에 모아상사_실습_2의 주간회의_2026-10-26.docx 내용 전체를 붙여넣기
==================================================
아래 회의 원문으로 파일을 만들어 줘.
1. 할 일 표: Excel 파일(xlsx). 열은 할 일 / 담당자 / 마감일. 파일 이름은 '할일표'.
2. 담당자별 안내문: 할 일을 맡은 사람마다 Word 파일(docx)을 하나씩 만들어.
   파일 이름은 '안내문_이름'. 안내문에는 그 사람이 맡은 할 일과 마감일만 넣어.
원문에 없는 내용은 추가하지 마. 마감일이 원문에 없으면 '미정'으로 적어.
시작하기 전에 어떤 순서로 할지 계획을 먼저 짧게 보여 주고 진행해.

(여기에 주간회의_2026-10-26.docx 내용을 붙여넣기)

--------------------------------------------------
[에이전트(Agent)와 작업] ④ 다르게 나오면: 파일 없이 글로만 답했을 때 - 같은 대화
--------------------------------------------------
위 내용을 요청한 파일로 만들어 줘

==================================================
[폴더 연결] ③ 목록과 개수 받기 - 모아상사_실습_2를 연결한 새 대화
==================================================
연결한 폴더에 있는 파일 목록을 표로 보여 줘.
열은 파일 이름 / 확장자로 하고, 맨 아래에 확장자별 개수와 전체 개수를 적어 줘.
파일은 아무것도 바꾸지 마.

==================================================
[폴더 연결] ④ 연결하지 않은 곳 묻기 - 같은 대화
==================================================
바탕화면에 있는 다른 파일 목록도 보여 줘.

==================================================
[파일 분류와 비교 요약] ① 파일 분류하기 (Manual) - 같은 대화
==================================================
연결한 폴더의 파일을 종류별 하위 폴더로 나눠 줘.
- 하위 폴더: 회의록 / 실적 / 영수증 / 제안서 / 기타
- 파일 이름은 바꾸지 말고, 지우지도 마.
- 옮기기 전에 어떤 파일을 어느 폴더로 옮길지 표로 먼저 보여 줘.
  내가 "진행해"라고 하면 옮겨.

--------------------------------------------------
[파일 분류와 비교 요약] ① 계획 표가 정답과 맞을 때 - 같은 대화
--------------------------------------------------
진행해

==================================================
[파일 분류와 비교 요약] ② 제안서 비교 요약하고 저장하기 (Auto) - 같은 대화
==================================================
제안서 폴더에 있는 제안서 3건을 비교해 줘.
- 표의 열: 거래처 / 월 임대료(원) / 계약 기간(개월) / 계약 기간 총액(원) / 설치까지 걸리는 기간 / 근거 파일 이름
- 계약 기간 총액은 월 임대료 × 계약 기간으로 계산해.
- 표 아래에 '월 임대료가 가장 낮은 곳'과 '계약 기간 총액이 가장 낮은 곳'을 따로 적어.
- 연결한 폴더 안에 '결과' 폴더를 만들고, 표를 '제안서_비교.xlsx'로 저장해.
- 원본 파일은 바꾸지 마.

==================================================
[원격 작업(모바일 앱)] ② 휴대폰에서 지시하기 - 휴대폰 앱에서 앞 섹션의 같은 작업
==================================================
결과 폴더의 '제안서_비교.xlsx'를 복사해서 '제안서_비교_설치순.xlsx'를 만들어 줘.
설치까지 걸리는 기간이 짧은 순서로 행을 정렬해.
원래 파일은 그대로 둬.
"""

# ---------------- 쓰기 ----------------
W.reset_dir(L)
W.reset_dir(T)
W.write_text(L / "실습_요청문.txt", PROMPTS)

# 실적 xlsx 5개
for m in (7, 8, 9):
    wb = W.new_wb()
    W.add_sheet(wb, f"{m}월 실적", M.month_header(m), M.month_rows(m), widths=[12, 12, 12, 12],
                money_cols=(2, 3, 4), first=True, title_row=M.month_title(m))
    W.save_wb(wb, F / f"2026년_{m:02d}월_실적.xlsx")
wb = W.new_wb()
H1 = list(range(1, 7))
W.add_sheet(wb, "상반기 실적", ["팀"] + [f"{m}월 실적" for m in H1] + ["상반기 합계"],
            [[t] + [M.MONTH_ACTUAL[m][t] for m in H1] + [sum(M.MONTH_ACTUAL[m][t] for m in H1)] for t in M.TEAMS],
            widths=[12] + [10] * 7, money_cols=tuple(range(2, 9)), first=True,
            title_row="모아상사 2026년 상반기 팀별 실적 요약 (단위: 만원, 실습용 가상 회사)")
W.save_wb(wb, F / "2026년_상반기_실적요약.xlsx")
wb = W.new_wb()
MS = list(range(1, 13))
W.add_sheet(wb, "팀별 목표", ["팀"] + [f"{m}월 목표" for m in MS] + ["연간 목표"],
            [[t] + [M.MONTH_TARGET[m][t] for m in MS] + [sum(M.MONTH_TARGET[m][t] for m in MS)] for t in M.TEAMS],
            widths=[12] + [9] * 13, money_cols=tuple(range(2, 15)), first=True,
            title_row="모아상사 2026년 팀별 월 목표 (단위: 만원, 실습용 가상 회사)")
W.save_wb(wb, F / "2026년_팀별_목표.xlsx")

# 회의록 docx 6개
for day, text in past_texts.items():
    doc = W.new_doc()
    W.doc_lines(doc, text)
    W.save_doc(doc, F / f"주간회의_{day}.docx")
doc = W.new_doc()
W.doc_lines(doc, M.MEETING_1026_TXT)  # 1교시 회의 원문 그대로
W.save_doc(doc, F / "주간회의_2026-10-26.docx")

# 비상연락망.docx
doc = W.new_doc()
doc.add_heading("모아상사 경영지원팀 비상연락망", level=0)
doc.add_paragraph("(실습용 가상 회사) 기준일 2026-10-01")
W.doc_table(doc, ["연락 순서", "이름", "직급", "내선 번호"], [(a, b, c, d) for a, b, c, d in M.CONTACTS])
doc.add_paragraph("비상시에는 연락 순서대로 전화하고, 연락이 닿지 않으면 다음 사람에게 넘깁니다.")
doc.add_paragraph("휴대전화 번호는 개인정보라 이 실습 파일에는 넣지 않았습니다.")
W.save_doc(doc, F / "비상연락망.docx")

# 사무실_자리배치도.pptx
prs = Presentation()
prs.slide_width, prs.slide_height = Cm(33.867), Cm(19.05)
s = prs.slides.add_slide(prs.slide_layouts[6])
tb = s.shapes.add_textbox(Cm(1.5), Cm(0.8), Cm(30), Cm(1.6))
tb.text_frame.text = "모아상사 경영지원팀 사무실 자리배치도 (실습용 가상 회사)"
for r in tb.text_frame.paragraphs[0].runs:
    r.font.size, r.font.bold = PPt(24), True
    W.set_ea_font(r)
pos = [(3, 4), (3, 9), (10, 9), (3, 13), (24, 4), (24, 13)]
for (label, where), (x, y) in zip(M.SEATS, pos):
    shp = s.shapes.add_shape(1, Cm(x), Cm(y), Cm(6.5), Cm(3))
    shp.text_frame.text = f"{label}\n({where})"
    for p in shp.text_frame.paragraphs:
        for r in p.runs:
            r.font.size = PPt(14)
            W.set_ea_font(r)
note = s.shapes.add_textbox(Cm(1.5), Cm(17), Cm(30), Cm(1))
note.text_frame.text = "3층 경영지원팀 자리. 회의는 3층 대회의실에서 함."
for r in note.text_frame.paragraphs[0].runs:
    r.font.size = PPt(12)
    W.set_ea_font(r)
prs.save(F / "사무실_자리배치도.pptx")

# 영수증 jpg 3장 + pdf 1장
for fname in ("영수증_택시_2026-10-15.jpg", "영수증_식대_2026-10-16.jpg", "영수증_퀵서비스_2026-10-20.jpg"):
    store, d, tm, items = M.RECEIPTS[fname]
    total = sum(a for _, a in items)
    lines = [("영 수 증", 30, True, "c"), (store, 20, False, "c"), ("-", 0, False, "l"),
             (f"일시\t{M.label(d)} {tm}", 17, False, "l"), ("-", 0, False, "l")]
    lines += [(f"{n}\t{a:,}원", 17, False, "l") for n, a in items]
    lines += [("-", 0, False, "l"), (f"합계\t{total:,}원", 21, True, "l"), ("결제\t법인카드", 17, False, "l"),
              ("-", 0, False, "l"), ("실습용 가상 영수증 (모아상사 실습)", 14, False, "c")]
    W.receipt_jpg(F / fname, lines)
W.write_bytes(F / "영수증_사무용품_2026-10-22.pdf", receipt_data)

# 제안서 pdf 3장
for name, data in proposal_data.items():
    W.write_bytes(F / f"제안서_{name}.pdf", data)

made = sorted(p.name for p in F.iterdir())
assert made == sorted(M.FILES_2), set(made) ^ set(M.FILES_2)
assert all(p.is_file() for p in F.iterdir())

# ---------------- 강사용 ----------------
wb = W.new_wb()
W.add_sheet(wb, "작업 맡기기 정답", ["파일", "들어 있어야 할 내용"], [
    ("할일표.xlsx", "할 일 4건. 아래 안내문 3개의 할 일·담당자·마감일을 합친 것과 같음"),
    ("안내문_이서연.docx", "연말 비품 재고 조사 정리 / 2026-11-06(금)"),
    ("안내문_박민수.docx", "11월 법인카드 사용 안내 메일 발송 / 2026-10-30(금), 송년회 장소 후보 3곳 조사 / 미정"),
    ("안내문_최우진.docx", "회의실 예약 규칙 개정안 초안 공유 / 2026-11-04(수)"),
    ("안내문 개수", "3개. 김지현 팀장은 맡은 할 일이 없어 안내문 없음"),
], widths=[22, 80], first=True)
W.add_sheet(wb, "파일 20개와 분류", ["번호", "파일 이름", "확장자", "분류 정답"],
            [(i, n, n.rsplit(".", 1)[1], cls[n]) for i, n in enumerate(M.FILES_2, start=1)],
            widths=[6, 34, 8, 10])
W.add_sheet(wb, "확장자·분류 개수", ["구분", "항목", "개수"],
            [("확장자", e, ext[e]) for e in ("docx", "xlsx", "pdf", "jpg", "pptx")] + [("확장자", "전체", 20)]
            + [("하위 폴더", k, cc[k]) for k in ("회의록", "실적", "영수증", "제안서", "기타")] + [("하위 폴더", "합계", 20)],
            widths=[12, 10, 8])
W.add_sheet(wb, "제안서 비교", ["거래처(가상)", "월 임대료(원)", "계약 기간(개월)", "계약 기간 총액(원)", "설치까지", "포함 사항", "근거 파일"],
            [(p[0], p[1], p[2], M.proposal_total(p), f"계약 후 {p[3]}영업일", p[4], f"제안서_{p[0]}.pdf") for p in M.PROPOSALS],
            widths=[14, 14, 14, 16, 16, 10, 24], money_cols=(2, 4),
            notes=["제안서는 모두 \"A3 컬러 복합기 1대 임대\" 제안, 금액은 부가세 포함",
                   "월 임대료가 가장 낮은 곳: 가람테크렌탈(162,000원)",
                   "계약 기간 총액이 가장 낮은 곳: 온새미사무기(5,016,000원). 가람테크렌탈은 총액으로는 가장 높음",
                   "원문 대조: 제안서_온새미사무기.pdf에서 \"설치\"를 찾으면 \"설치: 계약 후 3영업일\" 한 곳"])
W.add_sheet(wb, "설치순", ["순서", "거래처", "설치까지"],
            [(i, p[0], f"{p[3]}영업일") for i, p in enumerate(by_install, start=1)], widths=[6, 14, 12])
W.add_sheet(wb, "저장 결과 확인", ["확인할 것", "정답"], [
    ("모아상사_실습_2 안의 폴더", "6개: 회의록(6) / 실적(5) / 영수증(4) / 제안서(3) / 기타(2) / 결과(1)"),
    ("결과 폴더 (파일 분류와 비교 요약 ③)", "제안서_비교.xlsx 하나"),
    ("결과 폴더 (원격 작업 ③)", "2개: 제안서_비교.xlsx, 제안서_비교_설치순.xlsx"),
    ("다섯 폴더의 파일", "합치면 여전히 20개, 모아상사_실습_2 바로 아래 남은 파일 없음"),
    ("연결하지 않은 곳 묻기", "연결한 폴더만 볼 수 있다고 답하거나 바탕화면 접근을 묻는 창 → 창이 뜨면 거부"),
], widths=[34, 80])
W.save_wb(wb, T / "정답_2교시.xlsx")

# 시연용 참고 결과 (Claude가 만들 결과의 정답 모양)
wb = W.new_wb()
W.add_sheet(wb, "할일표", ["할 일", "담당자", "마감일"],
            [(t, n, M.label(d) if d else "미정") for t, n, d in M.TODO_1026], widths=[36, 10, 16], first=True)
W.save_wb(wb, EX / "할일표.xlsx")
for name in owners:
    doc = W.new_doc()
    doc.add_heading(f"{name} {M.RANK[name]} 할 일 안내", level=1)
    W.doc_table(doc, ["할 일", "마감일"], [(t, M.label(d) if d else "미정") for t, n, d in M.TODO_1026 if n == name])
    W.save_doc(doc, EX / f"안내문_{name}.docx")
cmp_header = ["거래처", "월 임대료(원)", "계약 기간(개월)", "계약 기간 총액(원)", "설치까지 걸리는 기간", "근거 파일 이름"]
for fname, order in (("제안서_비교.xlsx", M.PROPOSALS), ("제안서_비교_설치순.xlsx", by_install)):
    wb = W.new_wb()
    W.add_sheet(wb, "비교", cmp_header,
                [(p[0], p[1], p[2], M.proposal_total(p), f"계약 후 {p[3]}영업일", f"제안서_{p[0]}.pdf") for p in order],
                widths=[14, 14, 14, 16, 18, 26], money_cols=(2, 4), first=True,
                notes=["월 임대료가 가장 낮은 곳: 가람테크렌탈(162,000원)",
                       "계약 기간 총액이 가장 낮은 곳: 온새미사무기(5,016,000원)"])
    W.save_wb(wb, EX / "결과" / fname)

W.write_text(ROOT / "README.md", """# Claude 에이전트 - 실습 파일 (2교시)

교안: [2️⃣ Claude 에이전트](https://app.notion.com/p/19b435fcbd9d8294add901ed1609db7d)

## 학습자용 (수업 자료로 배포)

`학습자용\\모아상사_실습_2` 폴더를 통째로 압축해 배포합니다(압축을 풀면 바로 모아상사_실습_2 폴더가 나오게).
`실습_요청문.txt`는 실습 폴더 **밖**에 두었습니다. 폴더 안에 넣으면 폴더 연결 ③의 개수(20개)가 어긋납니다.

| 파일 | 쓰는 곳 | 내용 |
| --- | --- | --- |
| 모아상사_실습_2/ (하위 폴더 없이 파일 20개) | 에이전트와 작업 ~ 원격 작업 | docx 7 · xlsx 5 · pdf 4 · jpg 3 · pptx 1. 분류 정답 회의록 6 · 실적 5 · 영수증 4 · 제안서 3 · 기타 2 |
| 모아상사_실습_2/주간회의_2026-10-26.docx | 에이전트와 작업 ① | 1교시 회의 원문 그대로(한 줄에 한 문단) |
| 모아상사_실습_2/제안서_*.pdf (3개) | 파일 분류와 비교 요약 ② | 월 임대료·계약 기간·계약 기간 총액·설치(계약 후 N영업일)·토너 포함/별도. 글자로 된 PDF라 \"설치\"로 찾을 수 있음 |
| 모아상사_실습_2/2026년_08월·09월_실적.xlsx | 파일 분류 | 3교시 실적 파일 명세와 같은 숫자(9월 파일은 3교시 파일과 같은 내용) |
| 실습_요청문.txt | 전 섹션 | 교안의 요청문을 단계 순서대로 모은 것 |

## 강사용 (배포하지 않음)

| 파일 | 쓰는 곳 | 내용 |
| --- | --- | --- |
| 정답_2교시.xlsx | 전 섹션 | 작업 맡기기 정답(할 일 4건·안내문 3개), 파일 20개와 분류, 확장자·분류 개수, 제안서 비교(포함 사항까지), 설치순, 저장 결과 확인 |
| 예시_결과/ | 에이전트와 작업 ④ · 파일 분류와 비교 요약 ② · 원격 작업 ③ | 할일표.xlsx, 안내문_박민수·이서연·최우진.docx, 결과/제안서_비교.xlsx, 결과/제안서_비교_설치순.xlsx. Claude가 만들 결과의 정답 모양(시연·대조용) |

## 교안 명세에 내용이 없어 임의로 채운 파일

교안 📁에는 \"이름에 맞는 짧은 내용\"으로만 적혀 있어 아래는 만들면서 정했습니다. 실습 정답에는 쓰이지 않습니다.

- 주간회의_2026-09-21 ~ 10-19.docx (5개): 1교시 원문과 같은 머리 형식, 안건 2개씩(법인카드 정산, 대청소, 복합기 임대 제안서 요청·접수, 9월 실적 자료). \"설치\"라는 단어는 넣지 않음
- 2026년_07월_실적.xlsx, 2026년_상반기_실적요약.xlsx, 2026년_팀별_목표.xlsx의 숫자, 08월 파일의 8월 목표·7월 실적 (9월 목표, 8·9월 실적만 교안 숫자)
- 영수증 4장(택시 14,800원 · 식대 44,000원 · 퀵 18,000원 · 사무용품 42,700원), 상호는 \"실습용 가상 ○○\"
- 비상연락망.docx(내선 번호만), 사무실_자리배치도.pptx(한 장)
- 제안서 PDF의 머리·끝 문장(수신, 제안사, \"위 조건으로 제안드립니다\")
""")
print("ok", ROOT)
