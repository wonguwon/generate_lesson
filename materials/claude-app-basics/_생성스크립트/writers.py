"""실습 파일 공통 출력 도우미 (txt, xlsx, docx, pdf, jpg, pptx).

한글 글꼴은 Windows 맑은 고딕(C:/Windows/Fonts/malgun.ttf)을 씁니다.
"""
import io
import shutil
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

HEAD = Font(bold=True, color="FFFFFF")
HEAD_FILL = PatternFill("solid", fgColor="1F4E78")
WRAP = Alignment(wrap_text=True, vertical="top")
KFONT = "맑은 고딕"
FONT_DIR = Path("C:/Windows/Fonts")
MALGUN = FONT_DIR / "malgun.ttf"
MALGUN_BD = FONT_DIR / "malgunbd.ttf"


def need_font():
    if not MALGUN.exists():
        raise SystemExit(f"한글 글꼴이 없습니다: {MALGUN} (맑은 고딕이 있는 Windows에서 실행하세요)")


def reset_dir(path):
    """차시 폴더 안의 학습자용·강사용을 지우고 새로 만든다 (assert를 모두 통과한 뒤에만 부름)."""
    path = Path(path)
    if path.exists():
        shutil.rmtree(path)
    path.mkdir(parents=True)


# ---------------- txt ----------------
def write_text(path, text):
    """Windows 메모장에서 한글이 깨지지 않게 UTF-8(BOM), CRLF로 저장."""
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_bytes(text.replace("\r\n", "\n").replace("\n", "\r\n").encode("utf-8-sig"))


# ---------------- xlsx ----------------
def add_sheet(wb, title, header, rows, widths=None, money_cols=(), first=False, notes=(), title_row=None):
    """머리글 1행 + 데이터. title_row가 있으면 1행에 제목, 2행에 머리글."""
    ws = wb.active if first else wb.create_sheet()
    ws.title = title
    if title_row:
        ws.append([title_row])
        ws["A1"].font = Font(bold=True, size=12)
    ws.append(list(header))
    hr = ws.max_row
    for c in ws[hr]:
        c.font, c.fill, c.alignment = HEAD, HEAD_FILL, WRAP
    for r in rows:
        ws.append(list(r))
    for row in ws.iter_rows(min_row=hr + 1):
        for c in row:
            c.alignment = WRAP
            if c.column in money_cols and isinstance(c.value, (int, float)):
                c.number_format = "#,##0"
    widths = widths or [14] * len(header)
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = f"A{hr + 1}"
    if notes:
        ws.append([])
        for n in notes:
            ws.append([n])
    return ws


def new_wb():
    return Workbook()


def save_wb(wb, path):
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)


# ---------------- docx ----------------
def new_doc():
    doc = Document()
    for name in ("Normal", "Title", "Heading 1", "Heading 2", "List Bullet"):
        st = doc.styles[name]
        st.font.name = KFONT
        st.element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), KFONT)
    doc.styles["Normal"].font.size = Pt(10.5)
    return doc


def doc_lines(doc, text):
    """텍스트 블록을 줄마다 문단 하나로 그대로 넣는다 (내용을 덧붙이지 않음)."""
    for line in text.rstrip("\n").split("\n"):
        doc.add_paragraph(line)


def doc_table(doc, header, rows, right_cols=(), empty_rows=0):
    t = doc.add_table(rows=1, cols=len(header))
    t.style = "Table Grid"
    for i, h in enumerate(header):
        cell = t.rows[0].cells[i]
        cell.text = h
        cell.paragraphs[0].runs[0].bold = True
    for r in rows:
        cells = t.add_row().cells
        for i, v in enumerate(r):
            cells[i].text = str(v)
            if i in right_cols:
                cells[i].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.RIGHT
    for _ in range(empty_rows):
        t.add_row()
    return t


def grey(paragraph_or_run_text, doc):
    p = doc.add_paragraph()
    run = p.add_run(paragraph_or_run_text)
    run.font.color.rgb = RGBColor(0x80, 0x80, 0x80)
    return p


def save_doc(doc, path):
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(path)


def docx_text(path):
    d = Document(str(path))
    out = [p.text for p in d.paragraphs]
    for t in d.tables:
        for row in t.rows:
            out.append(" | ".join(c.text for c in row.cells))
    return "\n".join(out)


# ---------------- pdf (reportlab + 맑은 고딕) ----------------
def pdf_bytes(title, lines, footer=None):
    """A4 한 장짜리 PDF를 만들어 bytes로 돌려준다. lines: (글자, 크기, 굵게) 목록."""
    need_font()
    from reportlab.lib.pagesizes import A4
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    from reportlab.pdfgen import canvas

    if "Malgun" not in pdfmetrics.getRegisteredFontNames():
        pdfmetrics.registerFont(TTFont("Malgun", str(MALGUN)))
        pdfmetrics.registerFont(TTFont("MalgunBd", str(MALGUN_BD if MALGUN_BD.exists() else MALGUN)))
    buf = io.BytesIO()
    c = canvas.Canvas(buf, pagesize=A4, invariant=1)
    c.setTitle(title)
    w, h = A4
    y = h - 72
    for text, size, bold in lines:
        if text == "":
            y -= size * 0.8
            continue
        c.setFont("MalgunBd" if bold else "Malgun", size)
        c.drawString(64, y, text)
        y -= size * 1.7
    if footer:
        c.setFont("Malgun", 8)
        c.drawString(64, 48, footer)
    c.showPage()
    c.save()
    return buf.getvalue()


def pdf_text(data):
    from pypdf import PdfReader
    return "\n".join(p.extract_text() for p in PdfReader(io.BytesIO(data)).pages)


def write_bytes(path, data):
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_bytes(data)


# ---------------- jpg (Pillow + 맑은 고딕) ----------------
def receipt_jpg(path, lines, width=520):
    """영수증 사진처럼 보이는 세로 이미지. lines: (글자, 크기, 굵게, 정렬) 목록, 정렬은 l/c/r."""
    need_font()
    from PIL import Image, ImageDraw, ImageFont

    fonts = {}

    def f(size, bold):
        key = (size, bold)
        if key not in fonts:
            fonts[key] = ImageFont.truetype(str(MALGUN_BD if bold and MALGUN_BD.exists() else MALGUN), size)
        return fonts[key]

    height = 60 + sum(int(s * 1.6) if t != "-" else 18 for t, s, _, _ in lines)
    img = Image.new("RGB", (width, height), (250, 249, 244))
    d = ImageDraw.Draw(img)
    y = 30
    for text, size, bold, align in lines:
        if text == "-":
            d.line((30, y + 8, width - 30, y + 8), fill=(120, 120, 120), width=1)
            y += 18
            continue
        if "\t" in text:  # 왼쪽 글자 \t 오른쪽 글자
            left, right = text.split("\t")
            d.text((30, y), left, font=f(size, bold), fill=(30, 30, 30))
            rw = d.textlength(right, font=f(size, bold))
            d.text((width - 30 - rw, y), right, font=f(size, bold), fill=(30, 30, 30))
        else:
            tw = d.textlength(text, font=f(size, bold))
            x = {"l": 30, "c": (width - tw) / 2, "r": width - 30 - tw}[align]
            d.text((x, y), text, font=f(size, bold), fill=(30, 30, 30))
        y += int(size * 1.6)
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    img.save(path, "JPEG", quality=90)


# ---------------- pptx ----------------
def set_ea_font(run, name=KFONT):
    """python-pptx 글자에 한글(동아시아) 글꼴까지 지정."""
    run.font.name = name
    rpr = run._r.get_or_add_rPr()
    for tag in ("a:ea",):
        el = rpr.find(qn(tag))
        if el is None:
            el = rpr.makeelement(qn(tag), {})
            rpr.append(el)
        el.set("typeface", name)


def pptx_text(path):
    from pptx import Presentation
    out = []
    for i, s in enumerate(Presentation(str(path)).slides, start=1):
        for sh in s.shapes:
            if sh.has_text_frame:
                out.append(f"{i}: " + sh.text_frame.text.replace("\n", " / "))
            if getattr(sh, "has_table", False) and sh.has_table:
                for r in sh.table.rows:
                    out.append(f"{i}: " + " | ".join(c.text for c in r.cells))
            if getattr(sh, "has_chart", False) and sh.has_chart:
                for p in sh.chart.plots:
                    out.append(f"{i}: chart " + ",".join(str(c) for c in p.categories)
                               + " = " + ",".join(str(v) for v in p.series[0].values))
    return "\n".join(out)
