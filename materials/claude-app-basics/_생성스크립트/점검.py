"""실습 파일 점검. usage: py -3.14 -X utf8 점검.py <claude-app-basics 폴더>

1) make_01~05.py를 임시 폴더에 처음부터 다시 실행해, 지금 폴더의 파일과 내용이 같은지 비교한다
   (xlsx는 셀 값, docx·pptx는 글자, pdf는 추출한 글자, jpg는 픽셀, 나머지는 바이트).
2) 학습자용 실습 폴더의 파일 개수·이름이 교안 📁 명세와 같은지 확인한다.
3) 빈 폴더, zip 파일, BOM 없는 txt·md(SKILL.md 제외)가 없는지 확인한다.
"""
import hashlib
import shutil
import subprocess
import sys
import tempfile
from pathlib import Path

from openpyxl import load_workbook
from PIL import Image

import moa_data as M
import writers as W

ROOT = Path(sys.argv[1]).resolve()
HERE = Path(__file__).resolve().parent
LESSONS = ["01_Claude_시작하기", "02_Claude_에이전트", "03_문서_발표자료_웹페이지", "04_프로젝트와_스킬", "05_커넥터와_예약작업"]
problems = []


def content(p):
    s = p.suffix.lower()
    if s == ".xlsx":
        wb = load_workbook(p)
        return [(ws.title, [tuple(r) for r in ws.iter_rows(values_only=True)]) for ws in wb]
    if s == ".docx":
        return W.docx_text(p)
    if s == ".pptx":
        return W.pptx_text(p)
    if s == ".pdf":
        return W.pdf_text(p.read_bytes())
    if s == ".jpg":
        return hashlib.sha256(Image.open(p).tobytes()).hexdigest()
    return p.read_bytes()


# 1) 다시 만들어 비교
with tempfile.TemporaryDirectory() as tmp:
    for i in range(1, 6):
        r = subprocess.run([sys.executable, "-X", "utf8", str(HERE / f"make_0{i}.py"), tmp], cwd=HERE,
                           capture_output=True, text=True, encoding="utf-8")
        if r.returncode != 0:
            problems.append(f"make_0{i}.py 실패:\n{r.stderr}")
    for lesson in LESSONS:
        a, b = ROOT / lesson, Path(tmp) / lesson
        fa = sorted(p.relative_to(a) for p in a.rglob("*") if p.is_file())
        fb = sorted(p.relative_to(b) for p in b.rglob("*") if p.is_file())
        if fa != fb:
            problems.append(f"{lesson}: 파일 목록이 다름 {set(fa) ^ set(fb)}")
            continue
        for rel in fa:
            if content(a / rel) != content(b / rel):
                problems.append(f"{lesson}/{rel}: 다시 만든 결과와 내용이 다름")
        print(f"{lesson}: {len(fa)}개 파일, 다시 만든 결과와 같음")

# 2교시 폴더의 9월 실적 파일과 3교시 실적 파일은 같은 내용이어야 함
if content(ROOT / "02_Claude_에이전트/학습자용/모아상사_실습_2/2026년_09월_실적.xlsx") !=         content(ROOT / "03_문서_발표자료_웹페이지/학습자용/모아상사_실습_3/2026년_09월_실적.xlsx"):
    problems.append("2교시와 3교시의 2026년_09월_실적.xlsx 내용이 다름")
else:
    print("2교시·3교시 2026년_09월_실적.xlsx: 같은 내용")

# 2) 학습자용 실습 폴더 명세
SPEC = {
    "02_Claude_에이전트/학습자용/모아상사_실습_2": sorted(M.FILES_2),
    "03_문서_발표자료_웹페이지/학습자용/모아상사_실습_3": ["2026년_09월_실적.xlsx"],
    "04_프로젝트와_스킬/학습자용/모아상사_실습_4": sorted(["회의메모_2026-11-02.txt", "회의메모_2026-11-09.txt",
                                                    "프로젝트_지침.txt", "팀원_명단.txt", "주간업무_양식.docx"]),
    "05_커넥터와_예약작업/학습자용/모아상사_실습_5": sorted(
        [f"Drive에_올릴_파일/주간회의록_{d}.docx" for d in M.DRIVE_FILES]
        + [f"나중에_올릴_파일/주간회의록_{d}.docx" for d in M.LATER_FILES]),
}
for rel, want in SPEC.items():
    d = ROOT / rel
    got = sorted(str(p.relative_to(d)).replace("\\", "/") for p in d.rglob("*") if p.is_file())
    ok = got == sorted(want)
    print(f"{rel}: 파일 {len(got)}개 (명세 {len(want)}개) {'일치' if ok else '불일치'}")
    if not ok:
        problems.append(f"{rel}: 명세와 다름 {set(got) ^ set(want)}")
if (ROOT / "01_Claude_시작하기/학습자용/모아상사_실습_1").exists():
    problems.append("1교시에 모아상사_실습_1 폴더가 있음 (만들지 않기로 함)")
for lesson in LESSONS:
    if not (ROOT / lesson / "학습자용" / "실습_요청문.txt").is_file():
        problems.append(f"{lesson}: 실습_요청문.txt 없음")

# 3) 빈 폴더, zip, BOM
for p in ROOT.rglob("*"):
    if "_생성스크립트" in p.parts and "__pycache__" in p.parts:
        continue
    if p.is_dir() and not any(p.iterdir()):
        problems.append(f"빈 폴더: {p.relative_to(ROOT)}")
    if p.is_file() and p.suffix.lower() == ".zip":
        problems.append(f"zip 파일: {p.relative_to(ROOT)}")
    if p.is_file() and p.suffix.lower() in (".txt", ".md") and "_생성스크립트" not in p.parts and p.name != "SKILL.md":
        if not p.read_bytes().startswith(b"\xef\xbb\xbf"):
            problems.append(f"BOM 없음: {p.relative_to(ROOT)}")

shutil.rmtree(HERE / "__pycache__", ignore_errors=True)
print()
if problems:
    print("문제 있음:")
    for x in problems:
        print(" -", x)
    sys.exit(1)
print("모두 일치")
